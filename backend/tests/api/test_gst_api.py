"""GST Credit through the API: import, the six figures, the lists, decisions, export, reminders, plans, roles, tenant
isolation and the audit trail. Real Postgres, real routes; the GSTR-2B files are invented (valid check characters)."""

from __future__ import annotations

import io
import uuid
from datetime import date

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models.audit_log import AuditLog
from app.models.gst import GstDecision, GstMatch, GstReturn, GstReturnRecord
from app.services.gst import builder as b
from app.services.gst import vocabulary as v
from tests.security.plans import put_on_plan
from tests.services.gst.helpers import RECIPIENT, SUPPLIER_A, SUPPLIER_B, M, supplier_invoices
from tests.services.helm import factories

SEP = date(2026, 9, 1)


def book_invoice(db, tenant, *, number, supplier, gstin, on, taxable, rate=18, name="Sahyadri Packaging",
                 buyer_gstin=RECIPIENT, extra=None):
    line = b.item(taxable, rate)
    tax = line["cgst"] + line["sgst"]
    entities = {"vendor_gstin": gstin, "buyer_gstin": buyer_gstin, "subtotal": f"{taxable:.2f}",
                "cgst_amount": f"{line['cgst']:.2f}", "sgst_amount": f"{line['sgst']:.2f}", "tax_amount": f"{tax:.2f}",
                "invoice_date": on.isoformat(), **(extra or {})}
    return factories.invoice(db, workspace=tenant.workspace, organization_id=tenant.organization.id, number=number,
                             supplier=name, amount=taxable + tax, on=on, supplier_gstin=gstin, extra=entities)


def statement(extra_sections=None, *, period="2026-09"):
    return b.to_bytes(b.gstr2b(RECIPIENT, period, {
        "b2b": [
            supplier_invoices([
                b.document("INV-1001", date(2026, 9, 5), 1000),
                b.document("INV-1003", date(2026, 9, 12), 2100),          # the books say 2000
                b.document("INV-2001", date(2026, 9, 14), 500),            # the books say INV-2002: suggested
                b.document("NE-1", date(2026, 9, 15), 300, itc_available="N", reason="P"),
            ]),
            supplier_invoices([b.document("X-7", date(2026, 9, 16), 700)], gstin=SUPPLIER_B, name="Deccan Plastics"),
        ],
        "cdnr": [supplier_invoices([b.document("CN-1", date(2026, 9, 20), 100, note="C")], notes=True)],
        **(extra_sections or {}),
    }))


@pytest.fixture()
def world(client, db_session, tenant):
    put_on_plan(db_session, tenant.organization, "business")
    db = db_session
    docs = {
        "match": book_invoice(db, tenant, number="INV-1001", supplier="a", gstin=SUPPLIER_A, on=date(2026, 9, 5), taxable=1000),
        "waiting": book_invoice(db, tenant, number="INV-1002", supplier="a", gstin=SUPPLIER_A, on=date(2026, 9, 10), taxable=2000),
        "diff": book_invoice(db, tenant, number="INV-1003", supplier="a", gstin=SUPPLIER_A, on=date(2026, 9, 12), taxable=2000),
        "near": book_invoice(db, tenant, number="INV-2002", supplier="a", gstin=SUPPLIER_A, on=date(2026, 9, 14), taxable=500),
        "other_buyer": book_invoice(db, tenant, number="OB-1", supplier="a", gstin=SUPPLIER_A, on=date(2026, 9, 9), taxable=100,
                                    buyer_gstin=b.make_gstin("29", "AABCX1111Z")),
    }
    db.commit()
    return docs


def url(tenant, tail, workspace=None):
    return f"/api/v1/workspaces/{(workspace or tenant.workspace).id}/gst{tail}"


def upload(client, tenant, data, persona=None, name="GSTR2B_092026.json", workspace=None):
    persona = persona or tenant.contributor
    return client.post(url(tenant, "/returns", workspace), headers=persona.headers, files={"file": (name, data)})


def lines(client, tenant, persona=None, **params):
    response = client.get(url(tenant, "/records"), headers=(persona or tenant.viewer).headers, params=params)
    assert response.status_code == 200, response.text
    return response.json()


def test_an_import_reports_the_period_the_counts_and_the_problems_in_words(client, tenant, world):
    response = upload(client, tenant, statement())
    assert response.status_code == 200, response.text
    body = response.json()
    result = body["results"][0]
    assert (result["gstin"], result["period"], result["status"]) == (RECIPIENT, "2026-09", "IMPORTED")
    assert result["counts"] == {"b2b": 5, "cdnr": 1} and result["record_count"] == 6 and result["problems"] == []
    assert body["imported"] == 1 and body["unchanged"] == 0


def test_every_status_is_found_exactly_where_it_was_planted(client, tenant, world):
    upload(client, tenant, statement())
    found = {(m["number"], m["status"]) for m in lines(client, tenant)["items"]}
    assert found == {
        ("INV-1001", v.MATCHED), ("INV-1003", v.MATCHED_WITH_DIFFERENCES), ("INV-2001", v.SUGGESTED),
        ("NE-1", v.NOT_ELIGIBLE), ("X-7", v.NOT_IN_BOOKS), ("CN-1", v.NOTE), ("INV-1002", v.WAITING_FOR_SUPPLIER),
    }  # the invoice billed to another GSTIN is not in these books


def test_the_six_figures_are_sums_over_those_lines(client, tenant, world):
    upload(client, tenant, statement())
    body = client.get(url(tenant, "/summary"), headers=tenant.viewer.headers).json()
    figures = {f["key"]: f for f in body["figures"]}
    # claimable: 1001 (180) + 1003 at the lower figure, the books' (360) + credit note (-18)
    assert figures["claimable_now"]["amount_micros"] == (180 + 360 - 18) * M and figures["claimable_now"]["count"] == 2
    assert figures["waiting"]["amount_micros"] == 360 * M and figures["waiting"]["count"] == 1
    assert figures["waiting"]["suppliers"] == 1
    assert figures["differences"]["count"] == 1 and figures["differences"]["amount_micros"] == 18 * M
    assert figures["not_in_books"]["count"] == 1 and figures["not_in_books"]["amount_micros"] == 126 * M
    assert figures["not_eligible"]["count"] == 1 and figures["not_eligible"]["amount_micros"] == 54 * M
    for figure in body["figures"]:
        assert figure["meaning"] and figure["law"]["url"].startswith("https://")
    assert body["gstins"] == [RECIPIENT] and body["periods"] == ["2026-09"] and body["last_import"]["version"] == "1.0"
    assert "not tax advice" in body["disclaimer"]


def test_the_exact_same_file_again_is_a_no_op_and_a_changed_file_replaces_it_keeping_history(client, db_session, tenant, world):
    first = upload(client, tenant, statement()).json()
    again = upload(client, tenant, statement()).json()
    assert again["results"][0]["status"] == "UNCHANGED" and again["unchanged"] == 1
    assert again["results"][0]["return_id"] == first["results"][0]["return_id"]
    changed = statement({"b2ba": [supplier_invoices([b.document("INV-1001A", date(2026, 9, 5), 1000,
                                                                original=("INV-1001", date(2026, 9, 5)))])]})
    third = upload(client, tenant, changed).json()
    assert third["results"][0]["status"] == "REPLACED"
    returns = db_session.execute(select(GstReturn).where(GstReturn.workspace_id == tenant.workspace.id)
                                 .order_by(GstReturn.imported_at)).scalars().all()
    assert [r.status for r in returns] == ["SUPERSEDED", "ACTIVE"] and returns[1].supersedes_id == returns[0].id
    history = client.get(url(tenant, "/returns"), headers=tenant.viewer.headers, params={"include_replaced": True}).json()
    assert len(history["items"]) == 2
    assert len(client.get(url(tenant, "/returns"), headers=tenant.viewer.headers).json()["items"]) == 1


def test_a_bad_file_is_refused_in_words_and_changes_nothing(client, db_session, tenant, world):
    response = upload(client, tenant, b"this is not json", name="x.json")
    assert response.status_code == 422 and response.json()["detail"]["code"] == "NOT_JSON"
    assert db_session.execute(select(GstReturn)).first() is None


def test_a_person_confirms_a_suggestion_and_it_becomes_a_match(client, db_session, tenant, world):
    upload(client, tenant, statement())
    suggested = next(m for m in lines(client, tenant)["items"] if m["status"] == v.SUGGESTED)
    assert suggested["work_item_id"] == str(world["near"].id) and suggested["method"] == v.METHOD_SUGGESTED
    response = client.post(url(tenant, f"/records/{suggested['record_id']}/decision"), headers=tenant.contributor.headers,
                           json={"decision": "CONFIRM"})
    assert response.status_code == 200, response.text
    after = next(m for m in lines(client, tenant)["items"] if m["number"] == "INV-2001")
    assert (after["status"], after["method"], after["decision"]) == (v.MATCHED, v.METHOD_PERSON, "CONFIRM")
    audit = db_session.execute(select(AuditLog).where(AuditLog.resource_type == "GST_MATCH")).scalars().all()
    assert any((row.details or {}).get("decision") == "CONFIRM" for row in audit)
    undone = client.delete(url(tenant, f"/records/{suggested['record_id']}/decision"), headers=tenant.contributor.headers)
    assert undone.status_code == 200
    assert next(m for m in lines(client, tenant)["items"] if m["number"] == "INV-2001")["status"] == v.SUGGESTED


def test_not_ours_is_kept_across_a_re_import_and_leaves_the_not_in_books_count(client, db_session, tenant, world):
    upload(client, tenant, statement())
    record = next(m for m in lines(client, tenant)["items"] if m["number"] == "X-7")
    assert record["ims_action"] == v.IMS_PENDING
    done = client.post(url(tenant, f"/records/{record['record_id']}/decision"), headers=tenant.contributor.headers,
                       json={"decision": "NOT_OURS", "note": "Never bought from them"})
    assert done.status_code == 200
    row = next(m for m in lines(client, tenant)["items"] if m["number"] == "X-7")
    assert row["decision"] == "NOT_OURS" and row["ims_action"] == v.IMS_REJECT
    figures = {f["key"]: f for f in client.get(url(tenant, "/summary"), headers=tenant.viewer.headers).json()["figures"]}
    assert figures["not_in_books"]["count"] == 0
    changed = b.to_bytes(b.gstr2b(RECIPIENT, "2026-09", {"b2b": [supplier_invoices(
        [b.document("X-7", date(2026, 9, 16), 700)], gstin=SUPPLIER_B, name="Deccan Plastics")]}))
    upload(client, tenant, changed)
    assert next(m for m in lines(client, tenant)["items"] if m["number"] == "X-7")["decision"] == "NOT_OURS"


def test_a_link_by_hand_pairs_a_record_with_a_document(client, tenant, world):
    upload(client, tenant, statement())
    record = next(m for m in lines(client, tenant)["items"] if m["number"] == "X-7")
    bad = client.post(url(tenant, f"/records/{record['record_id']}/decision"), headers=tenant.contributor.headers,
                      json={"decision": "LINK", "work_item_id": str(uuid.uuid4())})
    assert bad.status_code == 404
    good = client.post(url(tenant, f"/records/{record['record_id']}/decision"), headers=tenant.contributor.headers,
                       json={"decision": "LINK", "work_item_id": str(world["waiting"].id)})
    assert good.status_code == 200, good.text
    linked = next(m for m in lines(client, tenant)["items"] if m["number"] == "X-7")
    assert linked["status"] in (v.MATCHED, v.MATCHED_WITH_DIFFERENCES) and linked["work_item_id"] == str(world["waiting"].id)
    nothing = client.post(url(tenant, f"/records/{record['record_id']}/decision"), headers=tenant.contributor.headers,
                          json={"decision": "CONFIRM"})
    assert nothing.status_code == 409 and nothing.json()["detail"]["code"] == "NOTHING_TO_CONFIRM"


def test_one_line_shows_both_sides_with_the_differences(client, tenant, world):
    upload(client, tenant, statement())
    row = next(m for m in lines(client, tenant)["items"] if m["number"] == "INV-1003")
    detail = client.get(url(tenant, f"/records/{row['id']}"), headers=tenant.viewer.headers).json()
    assert detail["books"]["taxable_micros"] == 2000 * M and detail["gstr2b"]["taxable_micros"] == 2100 * M
    assert detail["differences"][0]["field"] == "taxable" and detail["status_text"]
    assert detail["gstr2b"]["raw"]["inum"] == "INV-1003"


def test_filters_and_bounded_paging(client, tenant, world):
    upload(client, tenant, statement())
    assert lines(client, tenant, status="NOT_IN_BOOKS")["total"] == 1
    assert {m["number"] for m in lines(client, tenant, supplier="Deccan")["items"]} == {"X-7"}
    assert lines(client, tenant, ims="PENDING")["total"] >= 2
    page = lines(client, tenant, limit=2, offset=1)
    assert len(page["items"]) == 2 and page["total"] == 7 and page["offset"] == 1
    assert client.get(url(tenant, "/records"), headers=tenant.viewer.headers, params={"offset": 2**40}).status_code == 422
    assert client.get(url(tenant, "/records"), headers=tenant.viewer.headers, params={"limit": 999}).status_code == 422
    assert client.get(url(tenant, "/records"), headers=tenant.viewer.headers, params={"status": "NOPE"}).status_code == 422
    assert client.get(url(tenant, "/records"), headers=tenant.viewer.headers, params={"gstin": "12"}).status_code == 422


def test_suppliers_are_listed_oldest_waiting_first(client, tenant, world):
    upload(client, tenant, statement())
    body = client.get(url(tenant, "/suppliers"), headers=tenant.viewer.headers).json()
    first = body["items"][0]
    assert first["supplier_gstin"] == SUPPLIER_A and first["waiting_count"] == 1 and first["oldest_waiting"] == "2026-09-10"
    assert {s["supplier_gstin"] for s in body["items"]} == {SUPPLIER_A, SUPPLIER_B}


def test_the_export_is_a_workbook_with_the_six_sheets_and_is_audited(client, db_session, tenant, world):
    upload(client, tenant, statement())
    response = client.get(url(tenant, "/export.xlsx"), headers=tenant.contributor.headers, params={"gstin": RECIPIENT})
    assert response.status_code == 200 and "spreadsheetml" in response.headers["content-type"]
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Summary", "Matched", "Differences", "Waiting", "Not in books", "IMS actions"]
    waiting = [[c.value for c in row] for row in workbook["Waiting"].iter_rows(min_row=2)]
    assert waiting and waiting[0][2] == "INV-1002"
    assert any(row.action == "EXPORTED" for row in db_session.execute(select(AuditLog).where(AuditLog.resource_type == "GST_MATCH")).scalars())


def test_a_supplier_name_that_looks_like_a_formula_is_exported_as_text(client, tenant, world):
    hostile = b.to_bytes(b.gstr2b(RECIPIENT, "2026-09", {"b2b": [supplier_invoices(
        [b.document("H-1", date(2026, 9, 16), 700)], gstin=SUPPLIER_B, name="=HYPERLINK(\"http://x\",\"y\")")]}))
    upload(client, tenant, hostile)
    sheet = load_workbook(io.BytesIO(client.get(url(tenant, "/export.xlsx"), headers=tenant.contributor.headers).content))["Not in books"]
    assert sheet["A2"].value.startswith("'=")


def test_the_reminder_is_a_draft_listing_the_waiting_invoices_and_sends_nothing_without_a_known_contact(client, tenant, world):
    upload(client, tenant, statement())
    draft = client.get(url(tenant, f"/suppliers/{SUPPLIER_A}/reminder"), headers=tenant.contributor.headers,
                       params={"gstin": RECIPIENT}).json()
    assert [line["number"] for line in draft["lines"]] == ["INV-1002"] and "INV-1002" in draft["text"]
    assert draft["to"] is None and draft["can_send"] is False and "No email contact" in draft["note"]
    sent = client.post(url(tenant, f"/suppliers/{SUPPLIER_A}/reminder/send"), headers=tenant.contributor.headers,
                       params={"gstin": RECIPIENT}).json()
    assert sent["sent"] is False and "html" not in sent


def test_the_tolerance_is_a_workspace_setting_that_changes_the_result(client, tenant, world):
    upload(client, tenant, statement())
    assert client.patch(url(tenant, "/settings"), headers=tenant.contributor.headers, json={"tolerance_rupees": 500}).status_code == 403
    done = client.patch(url(tenant, "/settings"), headers=tenant.ws_admin.headers, json={"tolerance_rupees": 500})
    assert done.status_code == 200 and done.json()["tolerance_micros"] == 500 * M
    assert next(m for m in lines(client, tenant)["items"] if m["number"] == "INV-1003")["status"] == v.MATCHED


# ---------------------------------------------------------------------------------------- plans and roles


def test_a_developer_plan_sees_the_figures_and_counts_but_the_lists_are_locked(client, db_session, tenant, world):
    upload(client, tenant, statement())
    put_on_plan(db_session, tenant.organization, "developer")
    body = client.get(url(tenant, "/summary"), headers=tenant.viewer.headers).json()
    assert body["access"]["preview"] is True and "suppliers" in body["access"]["locked"]
    assert {f["key"]: f["count"] for f in body["figures"]}["waiting"] == 1
    for tail in ("/suppliers", "/records"):
        denied = client.get(url(tenant, tail), headers=tenant.viewer.headers)
        assert denied.status_code == 402, tail
    export = client.get(url(tenant, "/export.xlsx"), headers=tenant.contributor.headers)
    assert export.status_code == 402
    decision = client.post(url(tenant, f"/records/{uuid.uuid4()}/decision"), headers=tenant.contributor.headers,
                           json={"decision": "NOT_OURS"})
    assert decision.status_code == 402
    august = b.to_bytes(b.gstr2b(RECIPIENT, "2026-08", {}))
    assert upload(client, tenant, august).status_code == 200  # importing is part of the preview


def test_the_role_matrix(client, tenant, world):
    assert upload(client, tenant, statement(), persona=tenant.viewer).status_code == 403
    assert upload(client, tenant, statement(), persona=tenant.contributor).status_code == 200
    for persona in (tenant.viewer, tenant.contributor, tenant.ws_admin, tenant.owner, tenant.org_admin):
        assert client.get(url(tenant, "/summary"), headers=persona.headers).status_code == 200, persona
        assert client.get(url(tenant, "/records"), headers=persona.headers).status_code == 200
    assert client.get(url(tenant, "/export.xlsx"), headers=tenant.viewer.headers).status_code == 403
    assert client.get(url(tenant, "/export.xlsx"), headers=tenant.owner.headers).status_code == 200
    record = lines(client, tenant)["items"][0]
    assert client.post(url(tenant, f"/records/{record['record_id'] or uuid.uuid4()}/decision"), headers=tenant.viewer.headers,
                       json={"decision": "NOT_OURS"}).status_code == 403
    unauthenticated = client.get(url(tenant, "/summary"))
    assert unauthenticated.status_code in (401, 403)


def test_nothing_crosses_a_workspace_or_an_organization(client, db_session, tenant, world):
    upload(client, tenant, statement())
    record = lines(client, tenant)["items"][0]
    # another organization's member, on this workspace, and the other workspace with this one's data
    for persona in (tenant.other_org_member, tenant.non_member):
        for tail in ("/summary", "/records", "/suppliers", "/returns", "/export.xlsx"):
            assert client.get(url(tenant, tail), headers=persona.headers).status_code in (403, 404), (persona, tail)
        assert upload(client, tenant, statement(), persona=persona).status_code in (403, 404)
    # a member of another organization asking in their own workspace sees nothing of ours
    other = client.get(url(tenant, "/summary", tenant.foreign_workspace), headers=tenant.other_org_member.headers).json()
    assert other["gstins"] == [] and {f["count"] for f in other["figures"]} == {0}
    # ours: a record id of this workspace is a 404 in theirs
    assert client.get(url(tenant, f"/records/{record['id']}", tenant.foreign_workspace),
                      headers=tenant.other_org_member.headers).status_code in (402, 404)
    # and our people cannot reach their workspace
    assert client.get(url(tenant, "/summary", tenant.foreign_workspace), headers=tenant.owner.headers).status_code in (403, 404)
    assert db_session.execute(select(GstReturnRecord).where(GstReturnRecord.workspace_id == tenant.foreign_workspace.id)).first() is None


def test_an_import_is_audited(client, db_session, tenant, world):
    upload(client, tenant, statement())
    rows = db_session.execute(select(AuditLog).where(AuditLog.resource_type == "GST_RETURN")).scalars().all()
    assert len(rows) == 1 and rows[0].details["records"] == 6 and rows[0].details["gstin"] == RECIPIENT
    assert rows[0].actor_id == tenant.contributor.user.id
