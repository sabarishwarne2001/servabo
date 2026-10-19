"""The reconciliation as a workbook (openpyxl): Summary, Matched, Differences, Waiting, Not in books, IMS actions.

Text from documents and from suppliers is untrusted: a cell that would be read as a formula is written as text.
Amounts are rupees with two decimals, in plain numbers so a spreadsheet can add them."""

from __future__ import annotations

import io
import uuid
from datetime import date
from typing import Any, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.gst import GstMatch
from app.services.gst import summary
from app.services.gst import vocabulary as v

MAX_ROWS = 50_000
_INJECTION = ("=", "+", "-", "@", "\t", "\r")


def safe(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(_INJECTION):
        return "'" + value
    return value


def rupees(micros: Optional[int]) -> Optional[float]:
    return None if micros is None else round(micros / 1_000_000, 2)


def _sheet(workbook: Workbook, title: str, header: Sequence[str], rows: Sequence[Sequence[Any]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(list(header))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([safe(c) for c in row])
    for index, name in enumerate(header, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = max(12, min(40, len(name) + 4))


def build(db: Session, *, workspace_ids: Sequence[uuid.UUID], gstin: Optional[str], period: Optional[str],
          today: Optional[date] = None) -> bytes:
    today = today or date.today()
    figures = summary.figures(db, workspace_ids=workspace_ids, gstin=gstin, period=period, today=today)
    query = summary.scope(select(GstMatch), workspace_ids=workspace_ids, gstin=gstin, period=period)
    matches = list(db.execute(query.order_by(GstMatch.supplier_name, GstMatch.doc_date).limit(MAX_ROWS)).scalars())

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet.append(["GST Credit reconciliation"])
    summary_sheet["A1"].font = Font(bold=True, size=14)
    summary_sheet.append(["GSTIN", gstin or "all"])
    summary_sheet.append(["Period", period or "all imported"])
    summary_sheet.append(["As of", today.isoformat()])
    summary_sheet.append([])
    summary_sheet.append(["Figure", "Amount (INR)", "Count", "What it means"])
    for cell in summary_sheet[summary_sheet.max_row]:
        cell.font = Font(bold=True)
    for figure in figures["figures"]:
        summary_sheet.append([figure["title"], rupees(figure["amount_micros"]), figure["count"], figure["meaning"]])
    summary_sheet.append([])
    summary_sheet.append([v.NOT_ADVICE])
    summary_sheet.append([f"Law: {v.LAW_ITC_2B['section']}; {v.LAW_180_DAYS['section']}; {v.LAW_TIME_LIMIT['section']}"])
    summary_sheet.column_dimensions["A"].width = 34
    summary_sheet.column_dimensions["D"].width = 90

    def common(m: GstMatch) -> list[Any]:
        return [m.supplier_name, m.supplier_gstin, m.doc_number, m.doc_date.isoformat() if m.doc_date else None,
                rupees(m.value_micros), rupees(m.tax_micros), rupees(m.itc_micros)]

    head = ["Supplier", "Supplier GSTIN", "Number", "Date", "Invoice value", "GST", "Credit"]
    _sheet(workbook, "Matched", head + ["How matched", "Reported in"],
           [common(m) + [m.method, m.period_first_seen] for m in matches
            if m.status in (v.MATCHED, v.NOTE, v.SUGGESTED)])
    diff_rows = []
    for m in matches:
        if m.status == v.MATCHED_WITH_DIFFERENCES:
            for d in m.differences or []:
                diff_rows.append(common(m)[:4] + [d.get("label"), rupees(d.get("books_micros")),
                                                  rupees(d.get("gstr2b_micros")), rupees(d.get("delta_micros"))])
    _sheet(workbook, "Differences", ["Supplier", "Supplier GSTIN", "Number", "Date", "Field", "Your books",
                                     "GSTR-2B", "Books minus GSTR-2B"], diff_rows)
    _sheet(workbook, "Waiting", head + ["Section 16(4) last date", "Unpaid 180-day date"],
           [common(m) + [(m.risk or {}).get("deadline_16_4"), (m.risk or {}).get("reversal_date")]
            for m in matches if m.status == v.WAITING_FOR_SUPPLIER])
    _sheet(workbook, "Not in books", head + ["Reported in", "Decision"],
           [common(m) + [m.period_first_seen, m.decision] for m in matches if m.status == v.NOT_IN_BOOKS])
    _sheet(workbook, "IMS actions", head + ["Recommended action", "Why"],
           [common(m) + [m.ims_action, m.ims_reason] for m in matches if m.ims_action])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


__all__ = ["build", "rupees", "safe"]
